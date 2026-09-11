from datetime import date
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from services.dados_demo import validar_cliente, validar_mes, notas_do_mes, valor_total, estoque_cliente


def analisar_cliente(base, mes, ano):
    cliente = validar_cliente(base)
    mes, ano = validar_mes(mes, ano)
    notas = notas_do_mes(base, mes, ano)
    anterior_mes, anterior_ano = (12, ano - 1) if mes == 1 else (mes - 1, ano)
    anterior = notas_do_mes(base, anterior_mes, anterior_ano) if anterior_ano >= 2000 else []
    atual, passado = valor_total(notas), valor_total(anterior)
    estoque = estoque_cliente(base)
    situacao = "SEM FATURAMENTO" if atual == 0 else "CRESCIMENTO" if atual > passado else "QUEDA" if atual < passado else "ESTÁVEL"
    return {
        "sucesso": True, "demonstracao": True, "base": cliente["base"],
        "nome_empresa": cliente["nome"], "mes": mes, "ano": ano,
        "pedidos": len(notas), "faturamento_mes": atual,
        "faturamento_mes_anterior": passado, "variacao": round(atual - passado, 2),
        "situacao": situacao, "qtd_prod_controle_estoque": estoque["controle"],
        "qtd_prod_sob_demanda": estoque["sob_demanda"], "qtd_prod_flex": estoque["flex"],
        "qtd_prod_estoque_baixo": estoque["baixo"],
        "ultima_atualizacao_estoque": date(ano, mes, 1).isoformat() + " (simulada)",
        "ultima_venda": notas[-1]["data"] + " (simulada)" if notas else None,
    }


def gerar_planilha_analise(
    resultado: dict,
) -> BytesIO:

    # ============================================================
    # ARQUIVO EM MEMÓRIA
    # ============================================================

    arquivo = BytesIO()

    workbook = Workbook()

    sheet = workbook.active

    if sheet is None:

        raise RuntimeError(
            "Não foi possível criar a planilha."
        )

    sheet.title = (
        "Análise"
    )

    # ============================================================
    # TÍTULO
    # ============================================================

    sheet[
        "A1"
    ] = (
        "ANÁLISE DE CLIENTE — DADOS FICTÍCIOS"
    )

    sheet[
        "A1"
    ].font = Font(
        bold=True,
        size=16,
    )

    sheet.merge_cells(
        "A1:B1"
    )

    sheet[
        "A1"
    ].alignment = Alignment(
        horizontal="center",
        vertical="center",
    )

    # ============================================================
    # DADOS
    # ============================================================

    dados = [

        (
            "Base",
            resultado.get(
                "base",
                "",
            ),
        ),

        (
            "Nome da empresa",
            resultado.get(
                "nome_empresa",
                "",
            ),
        ),

        (
            "Mês",
            resultado.get(
                "mes",
                "",
            ),
        ),

        (
            "Ano",
            resultado.get(
                "ano",
                "",
            ),
        ),

        (
            "Pedidos no período",
            resultado.get(
                "pedidos",
                0,
            ),
        ),

        (
            "Faturamento do mês",
            resultado.get(
                "faturamento_mes",
                0,
            ),
        ),

        (
            "Faturamento mês anterior",
            resultado.get(
                "faturamento_mes_anterior",
                0,
            ),
        ),

        (
            "Variação em R$",
            resultado.get(
                "variacao",
                0,
            ),
        ),

        (
            "Situação",
            resultado.get(
                "situacao",
                "",
            ),
        ),

        (
            "Produtos controle de estoque",
            resultado.get(
                "qtd_prod_controle_estoque",
                0,
            ),
        ),

        (
            "Produtos sob demanda",
            resultado.get(
                "qtd_prod_sob_demanda",
                0,
            ),
        ),

        (
            "Produtos flex",
            resultado.get(
                "qtd_prod_flex",
                0,
            ),
        ),

        (
            "Produtos com estoque baixo",
            resultado.get(
                "qtd_prod_estoque_baixo",
                0,
            ),
        ),

        (
            "Última atualização de estoque",
            resultado.get(
                "ultima_atualizacao_estoque",
                "",
            ),
        ),

        (
            "Última venda",
            resultado.get(
                "ultima_venda",
                "",
            ),
        ),
    ]

    # ============================================================
    # ESCREVER DADOS
    # ============================================================

    linha_atual = 3

    for titulo, valor in dados:

        celula_titulo = (
            sheet.cell(
                row=linha_atual,
                column=1,
            )
        )

        celula_valor = (
            sheet.cell(
                row=linha_atual,
                column=2,
            )
        )

        celula_titulo.value = (
            titulo
        )

        celula_valor.value = (
            valor
        )

        celula_titulo.font = Font(
            bold=True
        )

        celula_titulo.alignment = Alignment(
            vertical="center"
        )

        celula_valor.alignment = Alignment(
            vertical="center"
        )

        linha_atual += 1

    # ============================================================
    # FORMATAÇÃO FINANCEIRA
    # ============================================================

    # Linha 8:
    # Faturamento do mês

    sheet[
        "B8"
    ].number_format = (
        'R$ #,##0.00'
    )

    # Linha 9:
    # Faturamento mês anterior

    sheet[
        "B9"
    ].number_format = (
        'R$ #,##0.00'
    )

    # Linha 10:
    # Variação

    sheet[
        "B10"
    ].number_format = (
        'R$ #,##0.00'
    )

    # ============================================================
    # LARGURA DAS COLUNAS
    # ============================================================

    sheet.column_dimensions[
        "A"
    ].width = 34

    sheet.column_dimensions[
        "B"
    ].width = 32

    # ============================================================
    # ALTURA DO TÍTULO
    # ============================================================

    sheet.row_dimensions[
        1
    ].height = 25

    # ============================================================
    # CONGELAR CABEÇALHO
    # ============================================================

    sheet.freeze_panes = (
        "A3"
    )

    # ============================================================
    # SALVAR EM MEMÓRIA
    # ============================================================

    workbook.save(
        arquivo
    )

    arquivo.seek(0)

    return arquivo
