from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from services.dados_demo import CLIENTES, validar_mes, notas_do_mes


def validar_periodo(mes_inicial, ano_inicial, mes_final, ano_final):
    mi, ai = validar_mes(mes_inicial, ano_inicial)
    mf, af = validar_mes(mes_final, ano_final)
    distancia = (af - ai) * 12 + mf - mi
    if distancia < 0:
        raise ValueError("O período final não pode ser anterior ao período inicial.")
    if distancia >= 36:
        raise ValueError("Selecione no máximo 36 meses por consulta.")
    return mi, ai, mf, af


def montar_meses_periodo(mes_inicial, ano_inicial, mes_final, ano_final):
    mi, ai, mf, af = validar_periodo(mes_inicial, ano_inicial, mes_final, ano_final)
    return [{"ano": valor // 12, "mes": valor % 12 + 1}
            for valor in range(ai * 12 + mi - 1, af * 12 + mf)]


def consultar_faturamento(mes_inicial, ano_inicial, mes_final, ano_final):
    mi, ai, mf, af = validar_periodo(mes_inicial, ano_inicial, mes_final, ano_final)
    meses = montar_meses_periodo(mi, ai, mf, af)
    detalhado, ranking, mensal = [], [], []
    por_mes = {(item["ano"], item["mes"]): [0, 0] for item in meses}
    total_centavos = total_pedidos = com_faturamento = 0
    for cliente in CLIENTES:
        centavos_cliente = pedidos_cliente = 0
        for periodo in meses:
            mes, ano = periodo["mes"], periodo["ano"]
            notas = notas_do_mes(cliente["base"], mes, ano)
            centavos = sum(nota["valor_centavos"] for nota in notas)
            pedidos = len(notas)
            centavos_cliente += centavos
            pedidos_cliente += pedidos
            por_mes[(ano, mes)][0] += centavos
            por_mes[(ano, mes)][1] += pedidos
            detalhado.append({"base": cliente["nome"], "mes": mes, "ano": ano,
                              "ocs_faturadas": pedidos, "faturamento": centavos / 100})
        total_centavos += centavos_cliente
        total_pedidos += pedidos_cliente
        com_faturamento += centavos_cliente > 0
        ranking.append({"base": cliente["nome"], "ocs_faturadas": pedidos_cliente,
                        "faturamento": centavos_cliente / 100})
    ranking.sort(key=lambda item: (-item["faturamento"], item["base"]))
    for posicao, item in enumerate(ranking, 1):
        item["posicao"] = posicao
    for periodo in meses:
        centavos, pedidos = por_mes[(periodo["ano"], periodo["mes"])]
        mensal.append({**periodo, "ocs_faturadas": pedidos, "faturamento": centavos / 100})
    return {
        "sucesso": True, "demonstracao": True,
        "mes_inicial": mi, "ano_inicial": ai, "mes_final": mf, "ano_final": af,
        "faturamento_total": total_centavos / 100, "ocs_faturadas": total_pedidos,
        "bases_consultadas": len(CLIENTES), "bases_com_faturamento": com_faturamento,
        "bases_sem_faturamento": len(CLIENTES) - com_faturamento, "bases_com_erro": 0,
        "faturamento_mensal": mensal, "ranking": ranking, "detalhado": detalhado, "erros": [],
    }


def gerar_planilha_faturamento(
    resultado: dict,
) -> BytesIO:
    """
    Gera a planilha completa do faturamento.

    Abas:
    1. Resumo Mensal
    2. Ranking Completo
    3. Detalhado por Cliente
    """

    arquivo = BytesIO()

    workbook = Workbook()

    # ============================================================
    # NOMES DOS MESES
    # ============================================================

    nomes_meses = [
        "",
        "Janeiro",
        "Fevereiro",
        "Março",
        "Abril",
        "Maio",
        "Junho",
        "Julho",
        "Agosto",
        "Setembro",
        "Outubro",
        "Novembro",
        "Dezembro",
    ]


    # ============================================================
    # ABA 1 - RESUMO MENSAL
    # ============================================================

    ws_resumo = workbook.active

    ws_resumo.title = (
        "Resumo Mensal"
    )


    ws_resumo["A1"] = (
        "FATURAMENTO — DADOS FICTÍCIOS"
    )

    ws_resumo["A1"].font = Font(
        bold=True,
        size=16,
    )


    periodo_inicial = (
        f"{nomes_meses[resultado['mes_inicial']]}/"
        f"{resultado['ano_inicial']}"
    )

    periodo_final = (
        f"{nomes_meses[resultado['mes_final']]}/"
        f"{resultado['ano_final']}"
    )


    ws_resumo["A2"] = (
        "Período"
    )

    ws_resumo["B2"] = (
        f"{periodo_inicial} a "
        f"{periodo_final}"
    )


    ws_resumo["A4"] = (
        "Faturamento total"
    )

    ws_resumo["B4"] = (
        resultado[
            "faturamento_total"
        ]
    )


    ws_resumo["A5"] = (
        "OCs faturadas"
    )

    ws_resumo["B5"] = (
        resultado[
            "ocs_faturadas"
        ]
    )


    ws_resumo["A6"] = (
        "Bases com faturamento"
    )

    ws_resumo["B6"] = (
        resultado[
            "bases_com_faturamento"
        ]
    )


    ws_resumo["A7"] = (
        "Bases sem faturamento"
    )

    ws_resumo["B7"] = (
        resultado[
            "bases_sem_faturamento"
        ]
    )


    ws_resumo["A8"] = (
        "Bases com erro"
    )

    ws_resumo["B8"] = (
        resultado[
            "bases_com_erro"
        ]
    )


    # ============================================================
    # CABEÇALHO FATURAMENTO MENSAL
    # ============================================================

    linha_inicio = 11

    ws_resumo.cell(
        row=linha_inicio,
        column=1,
        value="Mês",
    )

    ws_resumo.cell(
        row=linha_inicio,
        column=2,
        value="OCs faturadas",
    )

    ws_resumo.cell(
        row=linha_inicio,
        column=3,
        value="Faturamento",
    )


    for coluna in range(
        1,
        4,
    ):

        ws_resumo.cell(
            row=linha_inicio,
            column=coluna,
        ).font = Font(
            bold=True
        )


    # ============================================================
    # DADOS MENSAIS
    # ============================================================

    linha = (
        linha_inicio + 1
    )


    for item in resultado[
        "faturamento_mensal"
    ]:

        mes_texto = (
            f"{nomes_meses[item['mes']]}/"
            f"{item['ano']}"
        )


        ws_resumo.cell(
            row=linha,
            column=1,
            value=mes_texto,
        )

        ws_resumo.cell(
            row=linha,
            column=2,
            value=item[
                "ocs_faturadas"
            ],
        )

        ws_resumo.cell(
            row=linha,
            column=3,
            value=item[
                "faturamento"
            ],
        )


        ws_resumo.cell(
            row=linha,
            column=3,
        ).number_format = (
            'R$ #,##0.00'
        )


        linha += 1


    ws_resumo["B4"].number_format = (
        'R$ #,##0.00'
    )


    ws_resumo.column_dimensions[
        "A"
    ].width = 28

    ws_resumo.column_dimensions[
        "B"
    ].width = 22

    ws_resumo.column_dimensions[
        "C"
    ].width = 22


    # ============================================================
    # ABA 2 - RANKING COMPLETO
    # ============================================================

    ws_ranking = (
        workbook.create_sheet(
            "Ranking Completo"
        )
    )


    cabecalho_ranking = [
        "Posição",
        "Cliente/Base",
        "OCs faturadas",
        "Faturamento",
    ]


    for indice, titulo in enumerate(
        cabecalho_ranking,
        start=1,
    ):

        celula = ws_ranking.cell(
            row=1,
            column=indice,
            value=titulo,
        )

        celula.font = Font(
            bold=True
        )


    linha = 2


    for item in resultado[
        "ranking"
    ]:

        ws_ranking.cell(
            row=linha,
            column=1,
            value=item[
                "posicao"
            ],
        )

        ws_ranking.cell(
            row=linha,
            column=2,
            value=item[
                "base"
            ],
        )

        ws_ranking.cell(
            row=linha,
            column=3,
            value=item[
                "ocs_faturadas"
            ],
        )

        ws_ranking.cell(
            row=linha,
            column=4,
            value=item[
                "faturamento"
            ],
        )


        ws_ranking.cell(
            row=linha,
            column=4,
        ).number_format = (
            'R$ #,##0.00'
        )


        linha += 1


    ws_ranking.column_dimensions[
        "A"
    ].width = 12

    ws_ranking.column_dimensions[
        "B"
    ].width = 30

    ws_ranking.column_dimensions[
        "C"
    ].width = 20

    ws_ranking.column_dimensions[
        "D"
    ].width = 22


    # ============================================================
    # ABA 3 - DETALHADO POR CLIENTE
    # ============================================================

    ws_detalhado = (
        workbook.create_sheet(
            "Detalhado por Cliente"
        )
    )


    cabecalho_detalhado = [
        "Cliente/Base",
        "Mês",
        "OCs faturadas",
        "Faturamento",
    ]


    for indice, titulo in enumerate(
        cabecalho_detalhado,
        start=1,
    ):

        celula = ws_detalhado.cell(
            row=1,
            column=indice,
            value=titulo,
        )

        celula.font = Font(
            bold=True
        )


    linha = 2


    for item in resultado[
        "detalhado"
    ]:

        mes_texto = (
            f"{nomes_meses[item['mes']]}/"
            f"{item['ano']}"
        )


        ws_detalhado.cell(
            row=linha,
            column=1,
            value=item[
                "base"
            ],
        )

        ws_detalhado.cell(
            row=linha,
            column=2,
            value=mes_texto,
        )

        ws_detalhado.cell(
            row=linha,
            column=3,
            value=item[
                "ocs_faturadas"
            ],
        )

        ws_detalhado.cell(
            row=linha,
            column=4,
            value=item[
                "faturamento"
            ],
        )


        ws_detalhado.cell(
            row=linha,
            column=4,
        ).number_format = (
            'R$ #,##0.00'
        )


        linha += 1


    ws_detalhado.column_dimensions[
        "A"
    ].width = 30

    ws_detalhado.column_dimensions[
        "B"
    ].width = 20

    ws_detalhado.column_dimensions[
        "C"
    ].width = 20

    ws_detalhado.column_dimensions[
        "D"
    ].width = 22


    # ============================================================
    # ALINHAMENTOS
    # ============================================================

    for worksheet in [
        ws_resumo,
        ws_ranking,
        ws_detalhado,
    ]:

        for row in worksheet.iter_rows():

            for cell in row:

                cell.alignment = Alignment(
                    vertical="center"
                )


    # ============================================================
    # SALVAR NA MEMÓRIA
    # ============================================================

    workbook.save(
        arquivo
    )

    arquivo.seek(
        0
    )


    return arquivo
